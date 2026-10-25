#!/usr/bin/env python3
"""Rețeaua de localități a României prin filtrul Legii 351/2001 (anexa II, cu modificările Legii 100/2007):
pentru fiecare UAT (municipiu, oraș, comună) – indicatorii minimali de definire a localităților urbane, verificați
față de pragurile pentru oraș și pentru municipiu.

Intrări (descărcate automat în cache/urban/, dacă lipsesc):
  - RPL 2021, rezultate definitive (recensamantromania.ro): tabelul 1.22 (populația pe UAT și sate, cu cod SIRUTA),
    tabelul 5.30 (populația ocupată pe activități ale economiei naționale, pe UAT), tabelul 7.12 (locuințe pe UAT);
  - limitele UAT: Eurostat GISCO, LAU 2024 (1:1 milion), cod SIRUTA în GISCO_ID;
  - INS TEMPO-Online, matricele pe localități din data/urban/tempo/*.csv (scripts/urban_tempo.py, în GitHub Actions).
Ieșiri: public/data/urban/uat.geojson (limitele simplificate, doar la --geometrie, cu mapshaper prin npx)
        public/data/urban/urban_351.json (indicatorii și evaluarea, pe cod SIRUTA)
        public/data/urban/urban_351.csv (același conținut, pentru tabele)
"""
import argparse, csv, json, os, re, subprocess, sys, unicodedata, urllib.request

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
CACHE = os.path.join(ROOT, 'cache', 'urban')
TEMPO = os.path.join(ROOT, 'data', 'urban', 'tempo')
OUT = os.path.join(ROOT, 'public', 'data', 'urban')
CFG = json.load(open(os.path.join(ROOT, 'config', 'urban_351.json'), encoding='utf-8'))

RPL = 'https://www.recensamantromania.ro/wp-content/uploads/'
SURSE = {
    'Tabel-1.22.xlsx': RPL + '2023/05/Tabel-1.22.xlsx',
    'Tabel_5.30.xlsx': RPL + '2024/08/Tabel_5.30.xlsx',
    'RPL2021_Tabel-7.12.xlsx': RPL + '2026/02/RPL2021_Tabel-7.12.xlsx',
    'LAU_RG_01M_2024_4326.gpkg': 'https://gisco-services.ec.europa.eu/distribution/v2/lau/gpkg/LAU_RG_01M_2024_4326.gpkg',
}


def fisier(name):
    p = os.path.join(CACHE, name)
    if not os.path.exists(p):
        os.makedirs(CACHE, exist_ok=True)
        print('descarc', SURSE[name])
        urllib.request.urlretrieve(SURSE[name], p)
    return p


def fold(s):
    s = unicodedata.normalize('NFD', str(s)).encode('ascii', 'ignore').decode().upper()
    return re.sub(r'[^A-Z0-9]+', ' ', s).strip()


def fold_uat(s):
    """Denumirea UAT fără prefix (MUNICIPIUL / ORAȘ / COMUNA) și fără diacritice."""
    return re.sub(r'^(MUNICIPIUL|ORAS|COMUNA) ', '', fold(s))


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return None if v != v else float(v)
    t = str(v).strip().replace(' ', '')
    if t in ('', '-', ':', '*', 'x', '..'):
        return 0.0 if t == '-' else None
    try:
        return float(t.replace(',', '.'))
    except ValueError:
        return None


# ------------------------------------------------------------------ RPL 2021
def citeste_rpl():
    import openpyxl
    wb = openpyxl.load_workbook(fisier('Tabel-1.22.xlsx'), read_only=True)
    ws = wb.worksheets[0]
    import pyogrio.raw  # noqa: F401 – verificăm devreme dependența pentru limite
    meta, _f, _g, fields = pyogrio.raw.read(fisier('LAU_RG_01M_2024_4326.gpkg'), where="CNTR_CODE='RO'",
                                            read_geometry=False)
    coduri = {int(x.split('_')[1]) for x in fields[0]}
    uat, ordine, cur = {}, [], None
    for row in ws.iter_rows(min_row=6, values_only=True):
        jud, sir, nume, pop = row[0], row[1], row[2], row[3]
        if sir is None or nume is None:
            continue
        sir = int(sir)
        if sir in coduri:
            n = str(nume).strip()
            tip = 'M' if n.upper().startswith('MUNICIPIUL') else 'O' if re.match(r'^ORA[SȘŞ] ', n.upper()) else 'C'
            jn = str(jud).strip()
            if fold(jn) == 'MUNICIPIUL BUCURESTI':
                jn = 'București'
            cur = uat[sir] = {'siruta': sir, 'nume': n, 'judet': jn, 'tip': tip, 'pop2021': int(pop), 'sate': []}
            ordine.append(sir)
        elif cur is not None:
            cur['sate'].append([sir, str(nume).strip(), int(num(pop) or 0)])
    # 5.30: populația ocupată pe activități (fără cod SIRUTA – potrivire pe județ + denumire)
    wb = openpyxl.load_workbook(fisier('Tabel_5.30.xlsx'), read_only=True)
    ws = wb.worksheets[0]
    idx = {}
    for s in ordine:
        u = uat[s]
        idx[(fold(u['judet']), fold_uat(u['nume']))] = s
    jud = None
    gasite = 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        a, c = row[0], row[2]
        if a and str(a).strip().upper().startswith('JUDETUL'):
            jud = fold(str(a).strip()[8:])
            continue
        if a and fold(a) == 'MUNICIPIUL BUCURESTI':
            jud, c = 'BUCURESTI', 'MUNICIPIUL BUCURESTI'
        if not c or not jud:
            continue
        key = (jud, fold_uat(c))
        s = idx.get(key)
        if s is None:
            print('  5.30 fără corespondent:', key, file=sys.stderr)
            continue
        tot, agr = num(row[3]), num(row[4])
        if tot:
            uat[s]['ocupati'] = int(tot)
            uat[s]['ocupati_agricol'] = int(agr or 0)
            gasite += 1
    print('RPL 2021:', len(uat), 'UAT;', gasite, 'cu populația ocupată pe activități')
    return uat, ordine


# ------------------------------------------------------------------ limitele UAT
def geometrie(uat):
    import pyogrio.raw, shapely
    meta, _f, geom, fields = pyogrio.raw.read(fisier('LAU_RG_01M_2024_4326.gpkg'), where="CNTR_CODE='RO'")
    feats = []
    for i, g in enumerate(geom):
        s = int(fields[0][i].split('_')[1])
        feats.append({'type': 'Feature', 'properties': {'s': s},
                      'geometry': shapely.geometry.mapping(shapely.from_wkb(g))})
    raw = os.path.join(CACHE, 'uat_raw.geojson')
    json.dump({'type': 'FeatureCollection', 'features': feats}, open(raw, 'w'))
    os.makedirs(OUT, exist_ok=True)
    dest = os.path.join(OUT, 'uat.geojson')
    subprocess.run(['npx', '-y', 'mapshaper@0.6', raw, '-simplify', '12%', 'keep-shapes', '-clean',
                    '-o', 'format=geojson', 'precision=0.0001', dest], check=True)
    print('limite:', dest, os.path.getsize(dest) // 1024, 'KB')


def centre(uat):
    """Punctul reprezentativ al fiecărui UAT (în interiorul poligonului) – pentru simbolurile localităților."""
    import pyogrio.raw, shapely
    meta, _f, geom, fields = pyogrio.raw.read(fisier('LAU_RG_01M_2024_4326.gpkg'), where="CNTR_CODE='RO'")
    for i, g in enumerate(geom):
        s = int(fields[0][i].split('_')[1])
        if s in uat:
            p = shapely.from_wkb(g)
            c = p.centroid
            if not p.contains(c):
                c = p.representative_point()
            uat[s]['xy'] = [round(c.x, 4), round(c.y, 4)]
            uat[s]['km2'] = round(float(fields[5][i]), 1)


# ------------------------------------------------------------------ RPL 2021 pe județe (calibrarea estimărilor)
for _n, _u in [('RPL2021_Tabel-8.22.xlsx', '2026/02/RPL2021_Tabel-8.22.xlsx'), ('RPL2021_Tabel-8.9.xlsx', '2026/02/RPL2021_Tabel-8.9.xlsx'),
               ('RPL2021_Tabel-8.13.xlsx', '2026/02/RPL2021_Tabel-8.13.xlsx')]:
    SURSE[_n] = RPL + _u


def jud_key(s):
    s = fold(s)
    return 'BUCURESTI' if s in ('MUNICIPIUL BUCURESTI', 'BUCURESTI') else s


def rows(name):
    import openpyxl
    return list(openpyxl.load_workbook(fisier(name), read_only=True).worksheets[0].iter_rows(values_only=True))


def calibrare():
    """Totalurile pe județ din RPL 2021: locuințe, apă curentă în locuință (din rețea publică / sistem propriu),
    încălzire centrală (termoficare, centrală proprie pe gaze din rețea, alte centrale) pe județ × mediu,
    și raportul „baie și WC în locuință” / „apă curentă în locuință” pe medii (național)."""
    J = {}
    for r in rows('RPL2021_Tabel-8.22.xlsx'):
        if not r[0] or num(r[1]) is None:
            continue
        k = jud_key(r[0])
        J.setdefault(k, {}).update({'dw': num(r[1]), 'apa_loc': num(r[3]), 'apa_pub': num(r[4]), 'apa_prop': num(r[5]),
                                    'can_pub': num(r[28])})
    # 8.9: trei secțiuni (total, URBAN, RURAL), fiecare cu blocuri pe județe
    H, sect, cur = {}, 'T', None
    first_gas = False
    for r in rows('RPL2021_Tabel-8.9.xlsx'):
        t = str(r[0] or '').strip()
        if t in ('URBAN', 'RURAL'):
            sect = t[0]; cur = None; continue
        if t and t.isupper() and not t.startswith('-') and 'INCALZIRE' not in t and 'LOCUINTE' not in t and num(r[1]) is not None:
            cur = H.setdefault((jud_key(t), sect), {'dw': num(r[1])}); first_gas = True; continue
        if cur is None:
            continue
        tl = t.lower()
        if tl.startswith('cu incalzire centrala'):
            cur['central'] = num(r[1]) or 0
        elif tl.startswith('- termoficare'):
            cur['termo'] = num(r[1]) or 0
        elif tl.startswith('- gaze din retea publica') and first_gas:
            cur['gaz'] = num(r[1]) or 0; first_gas = False
    # 8.13: raportul baie+WC / apă în locuință, pe medii
    B, sect, last = {}, None, None
    for r in rows('RPL2021_Tabel-8.13.xlsx'):
        t = str(r[0] or '').strip()
        if t in ('ROMANIA', 'URBAN', 'RURAL'):
            sect = t; B[sect] = {}; continue
        if not sect:
            continue
        if t.startswith('a) Locuinte cu instalatie de alimentare cu apa'):
            last = 'apa'
        elif t.startswith('b) Locuinte cu baie'):
            last = 'baie'
        elif t.startswith('c) Locuinte cu closet'):
            last = 'wc'
        elif t.startswith('- in locuinta') and last and last not in B[sect]:
            B[sect][last] = num(r[1]); last = None
        elif t and not t.startswith('-'):
            last = None
    ratio = {m: min(B[m]['baie'], B[m]['wc']) / B[m]['apa'] for m in ('URBAN', 'RURAL')}
    print('calibrare RPL: %d județe, încălzire pe %d județ×mediu, baie+WC/apă: urban %.3f, rural %.3f'
          % (len(J), len(H), ratio['URBAN'], ratio['RURAL']))
    return J, H, ratio


def locuinte_uat(uat):
    """Locuințele convenționale pe UAT, RPL 2021 (tabelul 7.12; potrivire pe județ + denumire)."""
    idx = {(jud_key(u['judet']), fold_uat(u['nume'])): s for s, u in uat.items()}
    jud = None
    for r in rows('RPL2021_Tabel-7.12.xlsx'):
        t = str(r[0] or '').strip()
        if not t or num(r[3]) is None:
            continue
        k = (jud, fold_uat(t))
        if k in idx:
            uat[idx[k]]['dw2021'] = num(r[3])
        elif not re.match(r'^(MUNICIPIUL|ORA[SȘ]) ', t.upper()) or fold(t) == 'MUNICIPIUL BUCURESTI':
            jud = jud_key(t)
            if jud == 'BUCURESTI':
                s = [s for s, u in uat.items() if u['siruta'] == 179132][0]
                uat[s]['dw2021'] = num(r[3])
    lips = [u['nume'] for u in uat.values() if 'dw2021' not in u]
    print('locuințe RPL 2021 pe UAT:', len(uat) - len(lips), 'găsite', ('; lipsă: ' + ', '.join(lips[:8])) if lips else '')


# ------------------------------------------------------------------ INS TEMPO (data/urban/tempo/*.csv)
def tempo(code):
    """{siruta: {categorie: {an: valoare}}} din răspunsul brut al serverului (CSV cu antet)."""
    p = os.path.join(TEMPO, code + '.csv')
    out = {}
    if not os.path.exists(p):
        return None
    with open(p, encoding='utf-8') as f:
        rd = csv.reader(f)
        header = next(rd, None)
        for r in rd:
            r = [x.strip() for x in r]
            if len(r) < 3:
                continue
            sir = an = None
            cats = []
            for x in r[:-1]:
                m = re.match(r'^(\d+) [A-ZĂÂÎȘȚ]', x)
                y = re.match(r'^Anul (\d{4})$', x)
                if m and sir is None:
                    sir = int(m.group(1))
                elif y:
                    an = int(y.group(1))
                elif x and not x.startswith('UM') and x not in r[:1]:
                    cats.append(x)
            v = num(r[-1])
            if sir is None or an is None or v is None:
                continue
            # ultimele categorii sunt UM-ul (ex. „Kilometri”) – le păstrăm în cheie, nu deranjează
            out.setdefault(sir, {}).setdefault(' | '.join(cats), {})[an] = v
    return out


def ultim_an(T):
    ani = {}
    for cats in T.values():
        for ys in cats.values():
            for y, v in ys.items():
                ani[y] = ani.get(y, 0) + (1 if v else 0)
    # cel mai recent an cu date pentru cel puțin un sfert din localitățile cu date în anul cu cele mai multe
    if not ani:
        return None
    mx = max(ani.values())
    return max(y for y, n in ani.items() if n >= mx / 4)


def suma(T, s, an, test=lambda c: True):
    """Suma categoriilor care trec testul, pentru UAT-ul s în anul dat; 0 dacă UAT-ul lipsește din matrice
    (matricele INS pe localități listează doar localitățile cu unități / rețea), None dacă matricea lipsește."""
    if T is None or an is None:
        return None
    tot = 0.0
    for c, ys in T.get(s, {}).items():
        if test(c.lower()):
            tot += ys.get(an) or 0
    return tot


def titlu(s):
    mici = {'de', 'la', 'din', 'lui', 'pe', 'sub', 'cu', 'si', 'și'}
    s = re.sub(r'^(MUNICIPIUL|ORAȘ|ORAŞ|ORAS) ', '', str(s).strip())
    parts = re.split(r'([ \-])', s.lower())
    out = []
    for i, p in enumerate(parts):
        out.append(p if (p in mici and i > 0) else p[:1].upper() + p[1:])
    return ''.join(out)


def calculeaza():
    uat, ordine = citeste_rpl()
    locuinte_uat(uat)
    centre(uat)
    J, H, ratio = calibrare()
    T = {c: tempo(c) for c in ['POP107D', 'LOC101B', 'GOS104A', 'GOS105A', 'GOS106B', 'GOS110A', 'GOS103A', 'GOS108A',
                                'GOS109A', 'GOS116A', 'GOS118A', 'SAN102C', 'SAN104B', 'SCL101C', 'TUR102C', 'ART101B',
                                'ART103B', 'ART104A', 'SCL107C', 'SCL110B']}
    AN = {c: ultim_an(t) if t else None for c, t in T.items()}
    print('TEMPO – anii folosiți:', {c: a for c, a in AN.items()})
    seats = {k: v for k, v in CFG['resedinte_judet'].items()}
    rang1 = set(CFG['rang_I'])

    # consumurile pe UAT pentru estimări (apă casnică, gaze casnice, energie termică)
    def cons(code, test=lambda c: True):
        return {s: suma(T[code], s, AN[code], test) or 0 for s in uat} if T[code] else None
    apa = cons('GOS108A', lambda c: 'casnic' in c)
    gaz = cons('GOS118A', lambda c: 'casnic' in c)
    termo = cons('GOS109A')
    med = lambda u: 'U' if u['tip'] in 'MO' else 'R'

    # sumele pe județ (apă) și pe județ × mediu (încălzire), pentru alocarea totalurilor RPL 2021
    def grup(key):
        g = {}
        for s, u in uat.items():
            g.setdefault(key(u), []).append(s)
        return g
    GJ, GJM = grup(lambda u: jud_key(u['judet'])), grup(lambda u: (jud_key(u['judet']), med(u)))
    est = {}
    if apa:
        for j, ss in GJ.items():
            c = J.get(j)
            W = sum(apa[s] for s in ss)
            if not c or not W:
                continue
            P = {s: min(uat[s]['dw2021'], c['apa_pub'] * apa[s] / W) for s in ss}
            rest = sum(uat[s]['dw2021'] - P[s] for s in ss)
            for s in ss:
                dw = uat[s]['dw2021']
                S = (c['apa_loc'] - c['apa_pub']) * (dw - P[s]) / rest if rest > 0 else 0
                i3 = min(100.0, 100 * (P[s] + max(S, 0)) / dw) if dw else None
                est.setdefault(s, {})['i3'] = i3
                est[s]['i4'] = i3 * ratio['URBAN' if med(uat[s]) == 'U' else 'RURAL'] if i3 is not None else None
    if gaz and termo:
        for (j, m), ss in GJM.items():
            h = H.get((j, m)) or (H.get((j, 'T')) if j == 'BUCURESTI' else None)
            if not h or 'central' not in h:
                continue
            G, E = sum(gaz[s] for s in ss), sum(termo[s] for s in ss)
            Tu = {s: min(uat[s]['dw2021'], h.get('termo', 0) * termo[s] / E) if E else 0 for s in ss}
            Gu = {s: min(uat[s]['dw2021'] - Tu[s], h.get('gaz', 0) * gaz[s] / G) if G else 0 for s in ss}
            other = max(h['central'] - h.get('termo', 0) - h.get('gaz', 0), 0)
            rest = sum(uat[s]['dw2021'] - Tu[s] - Gu[s] for s in ss)
            for s in ss:
                dw = uat[s]['dw2021']
                O = other * (dw - Tu[s] - Gu[s]) / rest if rest > 0 else 0
                est.setdefault(s, {})['i5'] = min(100.0, 100 * (Tu[s] + Gu[s] + O) / dw) if dw else None

    out = []
    for s in ordine:
        u = uat[s]
        p = u['pop2021']
        v, x, an = {}, {}, {}
        # 2. populația ocupată în activități neagricole (RPL 2021)
        if u.get('ocupati'):
            v['i2'] = round(100 * (1 - u['ocupati_agricol'] / u['ocupati']), 1)
        # 3–5. dotarea locuințelor – estimări calibrate pe totalurile RPL 2021 ale județului
        for k in ('i3', 'i4', 'i5'):
            if est.get(s, {}).get(k) is not None:
                v[k] = round(est[s][k], 1); x[k] = 1
        # 6. paturi în spitale (internare continuă) ‰
        b = suma(T['SAN102C'], s, AN['SAN102C'], lambda c: 'spitale' in c and 'continua' in c)
        if b is not None:
            v['i6'] = round(1000 * b / p, 2); an['i6'] = AN['SAN102C']
        # 7. medici ‰ (toate formele de proprietate)
        m = suma(T['SAN104B'], s, AN['SAN104B'], lambda c: c.startswith('medici') and 'familie' not in c)
        if m is not None:
            v['i7'] = round(1000 * m / p, 2); an['i7'] = AN['SAN104B']
        # 8. învățământ: 2 = postliceal / universitar, 1 = liceal / profesional, 0 = fără
        if T['SCL101C']:
            post = suma(T['SCL101C'], s, AN['SCL101C'], lambda c: c.startswith('postliceal') or c.startswith('universitar'))
            lic = suma(T['SCL101C'], s, AN['SCL101C'], lambda c: c.startswith('liceal') or c.startswith('profesional'))
            v['i8'] = 2 if post else 1 if lic else 0; an['i8'] = AN['SCL101C']
        # 9. dotări culturale și sportive: 2 = bibliotecă publică + instituție de spectacole + bază sportivă,
        #    1 = bibliotecă publică + bază sportivă, 0 = altfel
        if T['ART101B'] and T['SCL107C']:
            lib = suma(T['ART101B'], s, AN['ART101B'], lambda c: c.startswith('publice'))
            spec = suma(T['ART103B'], s, AN['ART103B']) if T['ART103B'] else 0
            sport = (suma(T['SCL107C'], s, AN['SCL107C']) or 0) + (suma(T['SCL110B'], s, AN['SCL110B']) or 0 if T['SCL110B'] else 0)
            v['i9'] = 2 if lib and spec and sport else 1 if lib and sport else 0; an['i9'] = AN['ART101B']
            u['_cult'] = [int(lib or 0), int(spec or 0), int(sport or 0)]
        # 10. locuri în hoteluri
        h = suma(T['TUR102C'], s, AN['TUR102C'], lambda c: c.startswith('hoteluri') and 'tineret' not in c and 'apartament' not in c)
        if h is not None:
            v['i10'] = int(h); an['i10'] = AN['TUR102C']
        # 11–13. străzi (doar municipii și orașe: INS publică lungimea străzilor numai pentru urban)
        L = suma(T['GOS104A'], s, AN['GOS104A']) if u['tip'] in 'MO' else None
        if L:
            mod = suma(T['GOS105A'], s, AN['GOS105A'])
            if mod is not None:
                v['i11'] = round(min(100, 100 * mod / L), 1); an['i11'] = AN['GOS105A']
            ap = suma(T['GOS106B'], s, AN['GOS106B'])
            if ap is not None:
                v['i12'] = round(min(100, 100 * ap / L), 1); an['i12'] = AN['GOS106B']; x['i12'] = 1
            ca = suma(T['GOS110A'], s, AN['GOS110A'])
            if ca is not None:
                v['i13'] = round(min(100, 100 * ca / L), 1); an['i13'] = AN['GOS110A']; x['i13'] = 1
        # 16. spații verzi m²/locuitor (doar municipii și orașe)
        if u['tip'] in 'MO' and T['GOS103A']:
            g = suma(T['GOS103A'], s, AN['GOS103A'])
            v['i16'] = round(g * 10000 / p, 1); an['i16'] = AN['GOS103A']
        pd = suma(T['POP107D'], s, AN['POP107D']) if T['POP107D'] else None
        j = jud_key(u['judet'])
        nume = titlu(u['nume'])
        tip = u['tip']
        rj = 1 if fold(seats.get(j, '')) == fold(nume) and tip != 'C' else 0
        if s == 179132:
            r, rj = 0, 1
        elif tip == 'M':
            r = 1 if fold(nume) in {fold(x) for x in rang1} else 2
        else:
            r = 3 if tip == 'O' else 4
        sate = u['sate']
        rec = {'s': s, 'n': nume, 'j': titlu(u['judet']) if u['judet'] != 'București' else 'București', 't': tip, 'r': r,
               'p21': p, 'v': v}
        if pd:
            rec['pd'] = int(pd)
        if rj:
            rec['rj'] = 1
        if x:
            rec['x'] = x
        if an:
            rec['an'] = an
        if sate:
            rec['loc'] = titlu(sate[0][1]); rec['lp'] = sate[0][2]; rec['ns'] = len(sate)
            rec['sate'] = [[titlu(a[1]), a[2]] for a in sate]
        if u.get('xy'):
            rec['xy'] = u['xy']; rec['km2'] = u['km2']
        out.append(rec)
    return out, AN


def scrie(out, AN):
    ind = []
    surse = {
        'i1': 'INS, RPL 2021 (populația rezidentă) / TEMPO POP107D (după domiciliu)',
        'i2': 'INS, RPL 2021, tabelul 5.30 (populația ocupată pe activități ale economiei naționale)',
        'i3': 'Estimare: RPL 2021 (tabelul 8.22, pe județ) repartizat pe UAT după apa distribuită pentru uz casnic (TEMPO GOS108A)',
        'i4': 'Estimare: indicatorul 3 × raportul „baie și WC în locuință / apă în locuință” din RPL 2021 (tabelul 8.13, pe medii)',
        'i5': 'Estimare: RPL 2021 (tabelul 8.9, pe județ și mediu) repartizat după energia termică (GOS109A) și gazele casnice (GOS118A)',
        'i6': 'INS TEMPO SAN102C – paturi în spitale, internare continuă',
        'i7': 'INS TEMPO SAN104B – medici (toate formele de proprietate)',
        'i8': 'INS TEMPO SCL101C – unități de învățământ pe niveluri',
        'i9': 'INS TEMPO ART101B (biblioteci publice), ART103B (instituții de spectacole), SCL107C/SCL110B (săli de gimnastică și terenuri de sport școlare)',
        'i10': 'INS TEMPO TUR102C – capacitatea de cazare existentă în hoteluri',
        'i11': 'INS TEMPO GOS105A / GOS104A – străzi orășenești modernizate / lungimea străzilor',
        'i12': 'Aproximare: lungimea rețelei de apă (GOS106B) / lungimea străzilor (GOS104A), plafonată la 100%',
        'i13': 'Aproximare: lungimea canalizării (GOS110A) / lungimea străzilor (GOS104A), plafonată la 100%',
        'i14': 'Fără date publice pe localități (seria INS GOS111A se oprește în 2004)',
        'i15': 'Fără date publice pe localități',
        'i16': 'INS TEMPO GOS103A – suprafața spațiilor verzi (municipii și orașe)',
        'i17': 'Fără date publice pe localități',
    }
    scurt = {'i1': 'Locuitori', 'i2': 'Ocupați în activități neagricole', 'i3': 'Locuințe cu apă curentă',
             'i4': 'Locuințe cu baie și WC', 'i5': 'Locuințe cu încălzire centrală', 'i6': 'Paturi în spitale ‰',
             'i7': 'Medici ‰', 'i8': 'Învățământ', 'i9': 'Dotări culturale și sportive', 'i10': 'Locuri în hoteluri',
             'i11': 'Străzi modernizate', 'i12': 'Străzi cu rețea de apă', 'i13': 'Străzi cu canalizare',
             'i14': 'Epurarea apelor uzate', 'i15': 'Străzi cu hidranți', 'i16': 'Spații verzi m²/loc.', 'i17': 'Depozit de deșeuri'}
    coduri = {'i8': ['fără liceu', 'liceu / profesional', 'postliceal sau universitar'],
              'i9': ['incomplete', 'bibliotecă publică + bază sportivă', 'bibliotecă + instituție de spectacole + bază sportivă']}
    for i in CFG['indicatori']:
        d = dict(i)
        d['scurt'] = scurt[i['id']]
        d['sursa'] = surse[i['id']]
        if i['id'] in coduri:
            d['coduri'] = coduri[i['id']]; d['niveluri'] = True
        ind.append(d)
    pop_an = AN.get('POP107D')
    data = {'generat': __import__('datetime').date.today().isoformat(), 'an_domiciliu': pop_an, 'ani_tempo': AN,
            'indicatori': ind, 'uat': out,
            'info': 'Date: RPL 2021 (rezultate definitive) și INS TEMPO-Online, ultimii ani disponibili (%s). Generat la %s.'
                    % (', '.join(sorted({str(a) for a in AN.values() if a})), __import__('datetime').date.today().isoformat())}
    os.makedirs(OUT, exist_ok=True)
    with open(os.path.join(OUT, 'urban_351.json'), 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, separators=(',', ':'))
    with open(os.path.join(OUT, 'urban_351.csv'), 'w', encoding='utf-8', newline='') as f:
        w = csv.writer(f, delimiter=';')
        w.writerow(['siruta', 'uat', 'judet', 'tip', 'rang', 'pop_rpl2021', 'pop_domiciliu'] + ['ind%d' % i['nr'] for i in ind] + ['estimate'])
        for u in out:
            w.writerow([u['s'], u['n'], u['j'], u['t'], u['r'], u['p21'], u.get('pd', '')] +
                       [u['v'].get(i['id'], '') for i in ind] + [','.join(sorted(u.get('x', {})))])
    print('scris:', os.path.join(OUT, 'urban_351.json'), os.path.getsize(os.path.join(OUT, 'urban_351.json')) // 1024, 'KB')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--geometrie', action='store_true', help='regenerează limitele simplificate (mapshaper prin npx)')
    a = ap.parse_args()
    if a.geometrie:
        geometrie(None)
    out, AN = calculeaza()
    scrie(out, AN)


if __name__ == '__main__':
    main()
